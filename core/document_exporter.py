import os
import re
import html
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

class DocumentExporter:
    @staticmethod
    def _clean_option_text(opt_str):
        return re.sub(r'^(?:\([A-Da-d]\)|[A-Da-d][\.\)\:])\s*', '', str(opt_str or '').strip())

    @staticmethod
    def export_to_word(questions, filepath, title="MCQ Assessment", include_answers=True, course_meta=None, marks_per_question=2, duration="60 Mins"):
        doc = Document()
        
        try:
            mpq = float(marks_per_question) if marks_per_question and float(marks_per_question) > 0 else 2.0
        except (ValueError, TypeError):
            mpq = 2.0
        mpq_disp = int(mpq) if mpq.is_integer() else round(mpq, 1)
        tot_val = len(questions) * mpq
        tot_disp = int(tot_val) if tot_val.is_integer() else round(tot_val, 1)
        dur_str = str(duration or "60 Mins").strip()

        # Title
        heading = doc.add_heading(title, 0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Subtitle / Exam Metadata Header
        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mode_label = "FACULTY MASTER ANSWER KEY" if include_answers else "STUDENT EXAMINATION QUESTION PAPER"
        dept_str = (course_meta or {}).get("department", "Department of Computer Science & Engineering")
        code_str = (course_meta or {}).get("course_code", "CS301")
        sub_run = sub_p.add_run(
            f"{dept_str} • Course: {code_str}\n"
            f"{mode_label} • Duration: {dur_str} • Max Marks: {tot_disp} ({mpq_disp} Marks/Q)"
        )
        sub_run.bold = True
        sub_run.font.size = Pt(10)
        sub_run.font.color.rgb = RGBColor(79, 70, 229)

        if not include_answers:
            info_p = doc.add_paragraph()
            info_run = info_p.add_run("Student Name: ______________________________________    Roll No: ____________________    Date: ____________")
            info_run.bold = True
            info_run.font.size = Pt(10)

        doc.add_paragraph("")
        
        # Questions
        for i, q in enumerate(questions, 1):
            p = doc.add_paragraph()
            run_num = p.add_run(f"Q{i}. ")
            run_num.bold = True
            p.add_run(str(q.get('question', '')))
            
            for j, opt in enumerate(q.get('options', [])):
                opt_p = doc.add_paragraph()
                opt_p.paragraph_format.left_indent = Inches(0.4)
                letter_tag = chr(65 + j)
                opt_run = opt_p.add_run(f"({letter_tag}) ")
                opt_run.bold = True
                opt_p.add_run(DocumentExporter._clean_option_text(opt))
                
            if include_answers:
                ans_p = doc.add_paragraph()
                ans_p.paragraph_format.left_indent = Inches(0.4)
                clean_ans = DocumentExporter._clean_option_text(q.get('answer', ''))
                ans_run = ans_p.add_run(f"Correct Answer: {clean_ans}")
                ans_run.bold = True
                ans_run.font.color.rgb = RGBColor(34, 139, 34)
                if q.get("explanation"):
                    exp_p = doc.add_paragraph()
                    exp_p.paragraph_format.left_indent = Inches(0.4)
                    exp_label = exp_p.add_run("Rationale: ")
                    exp_label.bold = True
                    exp_label.font.size = Pt(9.5)
                    exp_label.font.color.rgb = RGBColor(79, 70, 229)
                    exp_body = exp_p.add_run(str(q.get("explanation", "")))
                    exp_body.font.size = Pt(9.5)
                    exp_body.font.color.rgb = RGBColor(71, 85, 105)
            
            doc.add_paragraph("") # Space between questions
            
        doc.save(filepath)

    @staticmethod
    def export_to_pdf(questions, filepath, title="MCQ Assessment", include_answers=True, course_meta=None, marks_per_question=2, duration="60 Mins"):
        doc = SimpleDocTemplate(
            filepath,
            pagesize=letter,
            rightMargin=45,
            leftMargin=45,
            topMargin=45,
            bottomMargin=45
        )
        
        try:
            mpq = float(marks_per_question) if marks_per_question and float(marks_per_question) > 0 else 2.0
        except (ValueError, TypeError):
            mpq = 2.0
        mpq_disp = int(mpq) if mpq.is_integer() else round(mpq, 1)
        tot_val = len(questions) * mpq
        tot_disp = int(tot_val) if tot_val.is_integer() else round(tot_val, 1)
        dur_str = html.escape(str(duration or "60 Mins").strip())

        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            alignment=1, # Center
            textColor=colors.HexColor('#1e1b4b'),
            spaceAfter=6
        )

        sub_style = ParagraphStyle(
            'SubStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=10,
            leading=14,
            alignment=1,
            textColor=colors.HexColor('#4f46e5'),
            spaceAfter=10
        )

        student_box_style = ParagraphStyle(
            'StudentBoxStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=10,
            leading=16,
            textColor=colors.HexColor('#1f2937'),
            spaceAfter=10
        )
        
        q_style = ParagraphStyle(
            'QuestionStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=16,
            textColor=colors.HexColor('#111827'),
            spaceBefore=8,
            spaceAfter=5
        )
        
        opt_style = ParagraphStyle(
            'OptionStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            leftIndent=20,
            textColor=colors.HexColor('#374151'),
            spaceAfter=3
        )
        
        ans_style = ParagraphStyle(
            'AnswerStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=9,
            leading=13,
            leftIndent=20,
            textColor=colors.HexColor('#15803d'),
            spaceBefore=3,
            spaceAfter=12
        )
        
        story = []
        safe_title = html.escape(str(title))
        story.append(Paragraph(safe_title, title_style))

        dept_str = html.escape(str((course_meta or {}).get("department", "Department of Computer Science & Engineering")))
        code_str = html.escape(str((course_meta or {}).get("course_code", "CS301")))
        mode_label = "FACULTY MASTER ANSWER KEY" if include_answers else "STUDENT EXAMINATION QUESTION PAPER"
        story.append(Paragraph(
            f"{dept_str} &bull; {code_str}<br/>{mode_label} &bull; Duration: {dur_str} &bull; Max Marks: {tot_disp} ({mpq_disp} Marks/Q)",
            sub_style
        ))

        if not include_answers:
            story.append(Paragraph("Student Name: ___________________________________ &nbsp;&nbsp;&nbsp; Roll No: ___________________ &nbsp;&nbsp;&nbsp; Date: ____________", student_box_style))

        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#e5e7eb'), spaceAfter=15))
        
        for i, q in enumerate(questions, 1):
            raw_q = html.escape(str(q.get('question', '')))
            q_text = f"<b>{i}.</b> {raw_q}"
            story.append(Paragraph(q_text, q_style))
            
            for j, opt in enumerate(q.get('options', [])):
                letter_tag = chr(65 + j)
                raw_opt = html.escape(DocumentExporter._clean_option_text(opt))
                opt_text = f"<b>({letter_tag})</b> {raw_opt}"
                story.append(Paragraph(opt_text, opt_style))
                
            if include_answers:
                raw_ans = html.escape(DocumentExporter._clean_option_text(q.get('answer', '')))
                raw_exp = html.escape(str(q.get('explanation', '')).strip())
                ans_text = f"<b>Answer:</b> {raw_ans}"
                if raw_exp:
                    ans_text += f"<br/><font color='#4f46e5'><b>Rationale:</b></font> <font color='#475569'>{raw_exp}</font>"
                story.append(Paragraph(ans_text, ans_style))
            story.append(Spacer(1, 6))
            
        doc.build(story)

    @staticmethod
    def export_class_report_pdf(faculty_details, stats, students, assignments, filepath):
        doc = SimpleDocTemplate(
            filepath,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=17,
            leading=22,
            textColor=colors.HexColor('#1e1b4b'),
            spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'ReportSub',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#4b5563'),
            spaceAfter=12
        )
        section_style = ParagraphStyle(
            'SectionHeader',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=12,
            leading=16,
            textColor=colors.HexColor('#312e81'),
            spaceBefore=12,
            spaceAfter=6
        )
        cell_style = ParagraphStyle(
            'CellText',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#1f2937')
        )
        cell_bold_style = ParagraphStyle(
            'CellBoldText',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#111827')
        )

        story = []
        fac_name = html.escape(str(faculty_details.get("name", "Faculty")))
        subject = html.escape(str(faculty_details.get("subject", "Course")))
        course_code = html.escape(str(faculty_details.get("course_code", "CS301")))
        dept = html.escape(str(faculty_details.get("department", "Department of CSE")))
        semester = html.escape(str(faculty_details.get("semester", "2026")))
        now_str = datetime.now().strftime("%d %b %Y, %I:%M %p")

        story.append(Paragraph(f"EduPulse Academic Performance & Attendance Report", title_style))
        story.append(Paragraph(f"<b>Faculty:</b> {fac_name} &nbsp;|&nbsp; <b>Course:</b> {course_code} - {subject} &nbsp;|&nbsp; <b>Department:</b> {dept} ({semester})<br/><b>Generated on:</b> {now_str}", subtitle_style))
        story.append(HRFlowable(width="100%", thickness=1.2, color=colors.HexColor('#c7d2fe'), spaceAfter=12))

        # 1. KPI Summary Table
        story.append(Paragraph("1. Class Cohort KPI Summary", section_style))
        kpi_data = [
            ["Total Enrolled", "Class Average Score", "Passing Rate (>=50%)", "Struggling (<50%)", "Mastery Index (>=75%)"],
            [
                str(len(students)),
                f"{stats.get('avg_score', '--')}%" if stats.get('avg_score') != '--' else '--',
                f"{stats.get('passing_rate', '--')}%" if stats.get('passing_rate') != '--' else '--',
                f"{stats.get('at_risk', 0)} Students",
                f"{stats.get('mastery', '--')}%" if stats.get('mastery') != '--' else '--'
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[100, 110, 115, 105, 110])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#eef2ff')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#312e81')),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 1), (-1, 1), 11),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ]))
        story.append(kpi_table)
        story.append(Spacer(1, 12))

        # 2. Tracked Assessments Breakdown Table
        story.append(Paragraph("2. Tracked Assessments Summary", section_style))
        asg_rows = [["Assessment Title", "Date", "Max Marks", "Submissions", "Class Avg", "Highest", "Lowest", "Pass Rate"]]
        for a in (assignments or []):
            if isinstance(a, dict):
                a_title = html.escape(str(a.get("title", "Assessment")))
                a_date = str(a.get("created_at", ""))[:10] or "--"
                m_val = float(a.get("max_marks", 100) or 100)
                a_max = str(int(m_val) if m_val.is_integer() else round(m_val, 1))
                a_sub = str(a.get("submission_count", 0))
                a_avg = f"{a.get('average', 0)}%"
                a_high = f"{a.get('highest', 0)}%"
                a_low = f"{a.get('lowest', 0)}%"
                a_pass = f"{a.get('pass_rate', 0)}%"
            else:
                a_title = html.escape(str(a[1] if len(a) > 1 else "Assessment"))
                a_date = str(a[2] if len(a) > 2 and a[2] else "")[:10] or "--"
                m_val = float(a[3] if len(a) > 3 and a[3] else 100)
                a_max = str(int(m_val) if m_val.is_integer() else round(m_val, 1))
                a_sub = "--"
                a_avg = "--"
                a_high = "--"
                a_low = "--"
                a_pass = "--"
            asg_rows.append([
                Paragraph(a_title, cell_bold_style),
                a_date,
                a_max,
                a_sub,
                a_avg,
                a_high,
                a_low,
                a_pass
            ])

        if len(asg_rows) == 1:
            asg_rows.append(["No assessments recorded yet", "--", "--", "--", "--", "--", "--", "--"])

        asg_table = Table(asg_rows, colWidths=[145, 65, 55, 60, 55, 52, 52, 56])
        asg_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#312e81')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ]))
        story.append(asg_table)
        story.append(Spacer(1, 12))

        # 3. Student Roster & Performance Table
        story.append(Paragraph("3. Student Performance & Attendance Roster", section_style))
        roster_rows = [["Roll No", "Student Name", "Attendance", "Prev Score", "Latest Score", "Overall Avg", "Status", "Diagnostic Insight"]]
        for s in students:
            prev_s = f"{s.get('prev')}%" if s.get('prev') not in (None, '--') else '--'
            lat_s = f"{s.get('latest')}%" if s.get('latest') not in (None, '--') else '--'
            avg_s = f"{s.get('overall_average')}%" if s.get('overall_average') not in (None, '--') else '--'
            safe_name = html.escape(str(s.get("name", "")))
            safe_insight = html.escape(str(s.get("insight", "")))
            roster_rows.append([
                str(s.get("roll", "")),
                Paragraph(safe_name, cell_bold_style),
                str(s.get("attendance", "--")),
                prev_s,
                lat_s,
                avg_s,
                str(s.get("status", "--")),
                Paragraph(safe_insight, cell_style)
            ])

        if len(roster_rows) == 1:
            roster_rows.append(["--", "No active students enrolled", "--", "--", "--", "--", "--", "--"])

        roster_table = Table(roster_rows, colWidths=[48, 98, 52, 48, 54, 52, 52, 136])
        roster_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e1b4b')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ALIGN', (2, 0), (6, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ]))
        story.append(roster_table)

        doc.build(story)

    @staticmethod
    def export_student_report_card_pdf(faculty_details, student_data, filepath):
        doc = SimpleDocTemplate(
            filepath,
            pagesize=letter,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'StudentReportTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#1e1b4b'),
            spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'StudentReportSub',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#4b5563'),
            spaceAfter=12
        )
        section_style = ParagraphStyle(
            'StudentSectionHeader',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=12,
            leading=16,
            textColor=colors.HexColor('#312e81'),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'StudentBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=15,
            textColor=colors.HexColor('#1f2937'),
            spaceAfter=8
        )

        story = []
        fac_name = html.escape(str(faculty_details.get("name", "Faculty")))
        subject = html.escape(str(faculty_details.get("subject", "Course")))
        course_code = html.escape(str(faculty_details.get("course_code", "CS301")))
        dept = html.escape(str(faculty_details.get("department", "Department of CSE")))
        semester = html.escape(str(faculty_details.get("semester", "2026")))
        now_str = datetime.now().strftime("%d %b %Y, %I:%M %p")

        s_name = html.escape(str(student_data.get("name", "Student")))
        s_roll = html.escape(str(student_data.get("roll", "--")))
        s_att = html.escape(str(student_data.get("attendance", "--")))
        s_grade = html.escape(str(student_data.get("grade", "--")))
        s_trend = html.escape(str(student_data.get("trend_status", "--")))
        s_insight = html.escape(str(student_data.get("insight", "No diagnostic remarks.")))

        story.append(Paragraph("EduPulse Individual Student Academic Report Card", title_style))
        story.append(Paragraph(
            f"<b>Course:</b> {course_code} - {subject} &nbsp;|&nbsp; <b>Faculty:</b> {fac_name} &nbsp;|&nbsp; <b>Department:</b> {dept} ({semester})<br/>"
            f"<b>Generated on:</b> {now_str}",
            subtitle_style
        ))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#4f46e5'), spaceAfter=14))

        # 1. Student Identity & Summary Metrics
        story.append(Paragraph("1. Student Profile & Performance Summary", section_style))
        total_taken = int(student_data.get("total_assessments", 0) or 0)
        try:
            att_num = float(str(student_data.get("attendance", "85")).replace("%", "").strip())
        except (ValueError, TypeError):
            att_num = 85.0
        att_display = f"{s_att} (Below 75% Min)" if att_num < 75.0 else s_att
        avg_display = f"{student_data.get('overall_average', 0)}%" if total_taken > 0 else "Not Assessed"
        best_display = f"{student_data.get('best_score', 0)}%" if total_taken > 0 else "--"
        low_display = f"{student_data.get('lowest_score', 0)}%" if total_taken > 0 else "--"
        traj_display = (
            f"{s_trend} ({'+' if float(student_data.get('growth_pts', 0)) >= 0 else ''}{student_data.get('growth_pts', 0)} pts)"
            if total_taken > 0 else "No Assessments Recorded"
        )
        summary_rows = [
            ["Student Name", s_name, "Roll Number", s_roll],
            ["Attendance Rate", att_display, "Academic Grade", s_grade],
            ["Overall Average", avg_display, "Trajectory Status", traj_display],
            ["Best Score", best_display, "Lowest Score", low_display]
        ]
        summary_table = Table(summary_rows, colWidths=[115, 150, 115, 150])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#eef2ff')),
            ('BACKGROUND', (2, 0), (2, -1), colors.HexColor('#eef2ff')),
            ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#312e81')),
            ('TEXTCOLOR', (2, 0), (2, -1), colors.HexColor('#312e81')),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTNAME', (2, 0), (2, -1), 'Helvetica-Bold'),
            ('FONTNAME', (1, 0), (1, -1), 'Helvetica-Bold'),
            ('FONTNAME', (3, 0), (3, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
            ('TOPPADDING', (0, 0), (-1, -1), 7),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 10))
        story.append(Paragraph(f"<b>Faculty Diagnostic Insight:</b> {s_insight}", body_style))

        # 2. Assessment History Table
        story.append(Paragraph("2. Chronological Assessment Breakdown", section_style))
        hist_rows = [["#", "Assessment Title", "Raw Score", "Normalized %", "Class Avg %", "vs. Class", "Submission Status"]]
        history = student_data.get("history", []) or []
        for idx, h in enumerate(history, 1):
            raw_sc = h.get("raw_score", h.get("score", 0))
            max_sc = h.get("max_score", 100)
            pct_sc = h.get("score", 0)
            cls_avg = h.get("class_average", 0)
            delta_cls = h.get("delta_from_class_avg", 0)
            delta_str = f"{'+' if float(delta_cls) >= 0 else ''}{delta_cls}%"
            hist_rows.append([
                str(idx),
                str(h.get("title", "Assessment")),
                f"{raw_sc} / {int(max_sc) if float(max_sc).is_integer() else max_sc}",
                f"{pct_sc}%",
                f"{cls_avg}%",
                delta_str,
                str(h.get("status", "On-time"))
            ])

        if len(hist_rows) == 1:
            hist_rows.append(["--", "No assessments recorded yet", "--", "--", "--", "--", "--"])

        hist_table = Table(hist_rows, colWidths=[28, 152, 72, 72, 68, 62, 76])
        hist_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e1b4b')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8.5),
            ('FONTSIZE', (0, 1), (-1, -1), 8.5),
            ('ALIGN', (0, 0), (0, -1), 'CENTER'),
            ('ALIGN', (2, 0), (-1, -1), 'CENTER'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
            ('TOPPADDING', (0, 0), (-1, -1), 7),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ]))
        story.append(hist_table)
        story.append(Spacer(1, 28))
        story.append(Paragraph(
            "<b>Faculty Signature:</b> _______________________________ &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; "
            "<b>Student / Guardian Signature:</b> _______________________________",
            body_style
        ))

        doc.build(story)

    @staticmethod
    def export_gradebook_csv(students, assignments, filepath):
        import csv
        # Order assignments chronologically
        ordered_asgs = sorted(assignments or [], key=lambda a: (str(a.get("created_at", "")), a.get("id", 0)))
        headers = ["Roll Number", "Student Name", "Attendance"]
        for a in ordered_asgs:
            max_m = a.get("max_marks", 100)
            max_label = int(max_m) if float(max_m).is_integer() else max_m
            headers.append(f"{a.get('title', 'Assessment')} (Out of {max_label})")
            headers.append(f"{a.get('title', 'Assessment')} (%)")
        headers.extend(["Overall Average (%)", "Best Score (%)", "Lowest Score (%)", "Submission Status", "Diagnostic Insight"])

        with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for s in students:
                marks_by_title = {}
                marks_by_id = {}
                for m in (s.get("all_marks") or []):
                    marks_by_title[m.get("title")] = m
                    if m.get("assignment_id") is not None:
                        marks_by_id[m.get("assignment_id")] = m

                row = [
                    s.get("roll", ""),
                    s.get("name", ""),
                    s.get("attendance", "--")
                ]
                for a in ordered_asgs:
                    m_obj = marks_by_id.get(a.get("id")) or marks_by_title.get(a.get("title"))
                    if m_obj:
                        row.append(m_obj.get("raw_score", m_obj.get("score", "")))
                        row.append(m_obj.get("score", ""))
                    else:
                        row.append("")
                        row.append("")
                row.extend([
                    s.get("overall_average", "--"),
                    s.get("best_score", "--"),
                    s.get("lowest_score", "--"),
                    s.get("status", "--"),
                    s.get("insight", "")
                ])
                writer.writerow(row)

    @staticmethod
    def export_sample_roster_csv(filepath):
        import csv
        sample_rows = [
            ["Student Name", "Roll Number", "Attendance %"],
            ["Aarav Sharma", "CS301-01", "92%"],
            ["Diya Patel", "CS301-02", "88%"],
            ["Rohan Kulkarni", "CS301-03", "95%"],
            ["Sneha Deshmukh", "CS301-04", "78%"],
            ["Vikram Joshi", "CS301-05", "85%"]
        ]
        with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerows(sample_rows)

