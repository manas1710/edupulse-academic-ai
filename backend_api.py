import json
import traceback
import os
if os.environ.get("EDUPULSE_WEB_MODE") == "1":
    webview = None
else:
    try:
        import webview
    except Exception:
        webview = None
from database.db_manager import DatabaseManager
from core.ai_handler import AIHandler
from core.document_exporter import DocumentExporter

class BackendAPI:
    def __init__(self):
        self.db = DatabaseManager()
        self.ai = AIHandler()
        
        # --- DEVELOPER INSTRUCTION ---
        # Configure your Gemini API key via the GEMINI_API_KEY environment variable or in Faculty Profile -> AI Settings.
        my_developer_key = ""
        saved_key = self.db.get_any_configured_api_key()
        active_key = saved_key or os.environ.get("GEMINI_API_KEY", "") or my_developer_key
        self.ai.set_api_key(active_key)

    def _get_export_dir(self):
        has_desktop_window = False
        if webview is not None:
            try:
                has_desktop_window = bool(webview.active_window() or len(webview.windows) > 0)
            except Exception:
                has_desktop_window = False
        desktop = os.path.join(os.environ.get('USERPROFILE', os.path.expanduser('~')), 'Desktop')
        if has_desktop_window and os.path.isdir(desktop):
            return desktop
        export_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data', 'exports')
        os.makedirs(export_dir, exist_ok=True)
        return export_dir

    def get_dashboard_stats(self, faculty_id=None):
        try:
            student_metrics = self.db.get_students_with_metrics(faculty_id)
            total_students = len(student_metrics)
            
            if total_students == 0:
                return {
                    "avg_score": "0",
                    "passing_rate": "0",
                    "at_risk": 0,
                    "mastery": "0",
                    "low_attendance_count": 0
                }
            
            low_att_count = 0
            for s in student_metrics:
                try:
                    att_val = float(str(s.get("attendance", "85%")).replace("%", "").strip())
                    if att_val < 75.0:
                        low_att_count += 1
                except ValueError:
                    pass

            # Get all marks for active students
            marks = self.db.get_all_marks(faculty_id)
            assessed_students = [
                s for s in student_metrics
                if s.get("latest_numeric") is not None and isinstance(s.get("overall_average"), (int, float))
            ]
            
            if not marks or not assessed_students:
                return {
                    "avg_score": "--",
                    "passing_rate": "--",
                    "at_risk": 0,
                    "mastery": "--",
                    "low_attendance_count": low_att_count
                }
            
            # Standardized cohort average & KPI thresholds across Dashboard and Student Progress:
            # Mastery >= 75%, Passing >= 50%, Struggling / At-Risk < 50%
            total_assessed = len(assessed_students)
            student_avgs = [float(s["overall_average"]) for s in assessed_students]
            avg = round(sum(student_avgs) / total_assessed, 1)
            
            passing_students = [s for s in assessed_students if float(s["overall_average"]) >= 50]
            passing_rate = round(len(passing_students) / total_assessed * 100, 1)
            
            at_risk = sum(1 for s in assessed_students if float(s["overall_average"]) < 50)
            
            mastery_students = [s for s in assessed_students if float(s["overall_average"]) >= 75]
            mastery = round(len(mastery_students) / total_assessed * 100, 1)
            
            return {
                "avg_score": str(avg),
                "passing_rate": str(passing_rate),
                "at_risk": at_risk,
                "mastery": str(mastery),
                "low_attendance_count": low_att_count
            }
        except Exception as e:
            print(f"Dashboard stats error: {e}")
            return {
                "avg_score": "--",
                "passing_rate": "--",
                "at_risk": 0,
                "mastery": "--",
                "low_attendance_count": 0
            }

    def download_dashboard_report(self, faculty_id=None):
        try:
            from datetime import datetime
            faculty_details = self.db.get_faculty_details(faculty_id) if faculty_id else None
            if not faculty_details:
                facs = self.db.get_all_faculties()
                first_id = facs[0][0] if facs else 1
                faculty_details = self.db.get_faculty_details(first_id) or {
                    "name": "Faculty",
                    "subject": "Data Science",
                    "course_code": "CS301",
                    "department": "Department of CSE",
                    "semester": "Semester V (2026)"
                }

            stats = self.get_dashboard_stats(faculty_id)
            students = self.db.get_students_with_metrics(faculty_id)
            progress_data = self.db.get_student_progress_analytics(faculty_id)
            assignments = progress_data.get("assignments", [])

            desktop = self._get_export_dir()
            course_code = "".join([c if c.isalnum() else "_" for c in str(faculty_details.get("course_code", "CS301"))])
            target_path = os.path.join(desktop, f"EduPulse_Class_Report_{course_code}.pdf")

            if os.path.exists(target_path):
                try:
                    with open(target_path, 'a'):
                        pass
                except (IOError, PermissionError):
                    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                    target_path = os.path.join(desktop, f"EduPulse_Class_Report_{course_code}_{ts}.pdf")

            DocumentExporter.export_class_report_pdf(faculty_details, stats, students, assignments, target_path)

            return {
                "status": "success",
                "path": target_path,
                "filename": os.path.basename(target_path),
                "download_url": f"/api/download/{os.path.basename(target_path)}"
            }
        except Exception as e:
            traceback.print_exc()
            return {"status": "error", "message": str(e)}

    def choose_file(self, filepath=None):
        try:
            if filepath and os.path.exists(filepath):
                return {"status": "success", "path": filepath, "name": os.path.basename(filepath)}
            if not webview:
                return {"status": "cancelled"}
            window = webview.active_window()
            if not window and len(webview.windows) > 0:
                window = webview.windows[0]
            if window:
                file_types = ('Document Files (*.pdf;*.docx;*.pptx;*.txt)', 'All files (*.*)')
                result = window.create_file_dialog(webview.OPEN_DIALOG, allow_multiple=False, file_types=file_types)
                if result:
                    path = result[0]
                    return {"status": "success", "path": path, "name": os.path.basename(path)}
            return {"status": "cancelled"}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def set_api_key(self, api_key, faculty_id=None):
        try:
            key = (api_key or "").strip()
            self.ai.set_api_key(key)
            os.environ["GEMINI_API_KEY"] = key
            if faculty_id is not None:
                self.db.update_faculty_api_key(faculty_id, key)
            return {"status": "success", "message": "API Key saved successfully."}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def _extract_text(self, filepath):
        ext = os.path.splitext(filepath)[1].lower()
        text = ""
        if ext == '.txt':
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
        elif ext == '.pdf':
            try:
                import PyPDF2
                with open(filepath, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    for page in reader.pages:
                        extracted = page.extract_text()
                        if extracted:
                            text += extracted + "\n"
            except Exception as pe:
                raise Exception(f"Failed to read PDF file: {str(pe)}")
        elif ext == '.docx':
            try:
                from docx import Document
                doc = Document(filepath)
                parts = []
                for p in doc.paragraphs:
                    if p.text and p.text.strip():
                        parts.append(p.text.strip())
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                        if row_cells:
                            parts.append(" | ".join(row_cells))
                text = "\n".join(parts)
            except Exception as pe:
                raise Exception(f"Failed to read Word (.docx) file: {str(pe)}")
        elif ext == '.pptx':
            try:
                from pptx import Presentation
                prs = Presentation(filepath)
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            text += shape.text + "\n"
            except Exception as pe:
                raise Exception(f"Failed to read PPTX file: {str(pe)}")
        else:
            raise Exception(f"Unsupported file format ({ext}). Please use PDF, DOCX, PPTX, or TXT.")
        return text

    def generate_mcqs(self, chapter, num_questions, difficulty, filepath=None, faculty_id=None, custom_text="", force_offline=False):
        try:
            if faculty_id is not None:
                fac_key = self.db.get_any_configured_api_key(faculty_id)
                if fac_key:
                    self.ai.set_api_key(fac_key)

            fac_details = self.db.get_faculty_details(faculty_id) if faculty_id else None

            syllabus_text = ""
            if custom_text and str(custom_text).strip():
                syllabus_text = str(custom_text).strip()
            else:
                # Resolve syllabus file path: explicit filepath -> faculty's saved syllabus_path -> dummy_syllabus.txt
                resolved_path = filepath if (filepath and os.path.exists(filepath)) else None
                if not resolved_path and fac_details:
                    saved_syl = fac_details.get("syllabus_path", "")
                    if saved_syl and os.path.exists(saved_syl):
                        resolved_path = saved_syl

                if resolved_path:
                    syllabus_text = self._extract_text(resolved_path)
                else:
                    dummy_path = os.path.join(os.path.dirname(__file__), "data", "dummy_syllabus.txt")
                    if os.path.exists(dummy_path):
                        syllabus_text = self._extract_text(dummy_path)

            if fac_details and fac_details.get("syllabus_units"):
                syllabus_text = (
                    f"Course: {fac_details.get('course_code', '')} - {fac_details.get('subject', '')}\n"
                    f"Syllabus Modules:\n{fac_details.get('syllabus_units', '')}\n\n"
                    + syllabus_text
                )
            
            # Ensure the text isn't empty
            if not syllabus_text.strip():
                return {"status": "error", "message": "The syllabus source appears to be empty or unreadable."}
                
            raw = self.ai.generate_mcqs(syllabus_text, chapter, int(num_questions), difficulty, force_offline=bool(force_offline))
            
            if isinstance(raw, str):
                questions = json.loads(raw)
            else:
                questions = raw
                
            return {
                "status": "success",
                "data": questions,
                "mode": getattr(self.ai, "last_generation_mode", "ai")
            }
        except Exception as e:
            traceback.print_exc()
            return {"status": "error", "message": str(e)}

    def export_mcqs(self, questions, file_format, topic=None, include_answers=True, faculty_id=None, marks_per_question=2, duration="60 Mins"):
        try:
            if not questions or len(questions) == 0:
                return {"status": "error", "message": "No questions selected for export."}

            desktop = self._get_export_dir()
            
            # Clean title and filename
            doc_title = topic.strip() if topic and topic.strip() else "MCQ Assessment"
            safe_name = "".join([c if c.isalnum() or c in " _-" else "_" for c in doc_title]).strip()
            if not safe_name:
                safe_name = "Generated_MCQs"
            
            mode_suffix = "_Answer_Key" if include_answers else "_Question_Paper"
            ext = ".docx" if file_format == "word" else ".pdf"
            target_path = os.path.join(desktop, f"{safe_name}{mode_suffix}{ext}")

            # Check if file exists and is locked (e.g. currently open in Word or Acrobat)
            if os.path.exists(target_path):
                try:
                    with open(target_path, 'a'):
                        pass
                except (IOError, PermissionError):
                    from datetime import datetime
                    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                    target_path = os.path.join(desktop, f"{safe_name}{mode_suffix}_{ts}{ext}")

            course_meta = self.db.get_faculty_details(faculty_id) if faculty_id else None

            if file_format == "word":
                DocumentExporter.export_to_word(
                    questions, target_path, title=doc_title,
                    include_answers=bool(include_answers), course_meta=course_meta,
                    marks_per_question=marks_per_question, duration=duration
                )
            elif file_format == "pdf":
                DocumentExporter.export_to_pdf(
                    questions, target_path, title=doc_title,
                    include_answers=bool(include_answers), course_meta=course_meta,
                    marks_per_question=marks_per_question, duration=duration
                )
            else:
                return {"status": "error", "message": f"Unsupported format: {file_format}"}

            return {
                "status": "success",
                "path": target_path,
                "filename": os.path.basename(target_path),
                "download_url": f"/api/download/{os.path.basename(target_path)}"
            }
        except Exception as e:
            traceback.print_exc()
            return {"status": "error", "message": str(e)}

    def open_file(self, file_path):
        try:
            if file_path and os.path.exists(file_path):
                os.startfile(file_path)
                return {"status": "success"}
            return {"status": "error", "message": "File not found"}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def open_folder(self, file_path):
        try:
            if file_path and os.path.exists(file_path):
                import subprocess
                subprocess.Popen(f'explorer /select,"{os.path.abspath(file_path)}"')
                return {"status": "success"}
            return {"status": "error", "message": "Folder not found"}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def get_faculties(self):
        faculties = self.db.get_all_faculties()
        return [{"id": f[0], "name": f[1], "subject": f[2], "profile_photo": f[3] if len(f) > 3 and f[3] else ""} for f in faculties]

    def get_faculty_profile(self, faculty_id):
        try:
            profile = self.db.get_faculty_details(faculty_id)
            if not profile:
                return {"status": "error", "message": "Faculty not found"}
            
            students = self.db.get_students(faculty_id)
            total_students = len(students)
            assignments = self.db.get_assignments(faculty_id)
            total_assignments = len(assignments)
            
            dash_stats = self.get_dashboard_stats(faculty_id)
            avg_score = dash_stats.get("avg_score", "--")
            
            student_metrics = self.db.get_students_with_metrics(faculty_id)
            att_values = []
            for sm in student_metrics:
                att_str = sm.get("attendance", "").replace("%", "").strip()
                try:
                    att_values.append(float(att_str))
                except (ValueError, TypeError):
                    pass
            avg_att = f"{round(sum(att_values)/len(att_values), 1)}%" if att_values else "--"
            
            profile["total_students"] = total_students
            profile["total_assignments"] = total_assignments
            profile["avg_score"] = avg_score
            profile["avg_attendance"] = avg_att
            raw_key = (profile.get("api_key") or "").strip()
            profile["has_api_key"] = bool(raw_key)
            profile["api_key_masked"] = (raw_key[:6] + "..." + raw_key[-4:]) if len(raw_key) > 10 else ("Configured" if raw_key else "")
            
            return {"status": "success", "data": profile}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def update_faculty_profile(self, faculty_id, data):
        try:
            import shutil
            raw_syl_path = (data or {}).get("syllabus_path", "")
            if raw_syl_path and os.path.exists(raw_syl_path):
                syllabi_dir = os.path.join(os.path.dirname(__file__), "data", "syllabi")
                os.makedirs(syllabi_dir, exist_ok=True)
                safe_base = "".join([c if c.isalnum() or c in "._- " else "_" for c in os.path.basename(raw_syl_path)])
                dest_path = os.path.join(syllabi_dir, f"faculty_{faculty_id}_{safe_base}")
                if os.path.abspath(raw_syl_path) != os.path.abspath(dest_path):
                    shutil.copy2(raw_syl_path, dest_path)
                data["syllabus_path"] = dest_path
            self.db.update_faculty_details(faculty_id, data)
            if (data or {}).get("api_key"):
                self.ai.set_api_key(str(data.get("api_key")).strip())
            return {"status": "success", "message": "Profile updated successfully", "syllabus_path": data.get("syllabus_path", "")}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def change_faculty_password(self, faculty_id, current_password, new_password):
        try:
            if not new_password or len(new_password.strip()) < 4:
                return {"status": "error", "message": "New password must be at least 4 characters."}
            stored_hash = self.db.get_faculty_password(faculty_id)
            if not stored_hash:
                return {"status": "error", "message": "Faculty profile not found."}
            curr_clean = (current_password or "").strip()
            if self._hash_password(current_password or "") != stored_hash and self._hash_password(curr_clean) != stored_hash:
                return {"status": "error", "message": "Current password is incorrect."}
            new_hash = self._hash_password(new_password.strip())
            self.db.update_faculty_password(faculty_id, new_hash)
            return {"status": "success", "message": "Password updated successfully."}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def save_faculty_api_key(self, faculty_id, api_key):
        try:
            clean_key = (api_key or "").strip()
            self.db.update_faculty_api_key(faculty_id, clean_key)
            if clean_key:
                self.ai.set_api_key(clean_key)
                os.environ["GEMINI_API_KEY"] = clean_key
            return {"status": "success", "message": "Gemini API Key saved to profile."}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def upload_faculty_photo(self, faculty_id, filepath=None):
        try:
            import base64
            import mimetypes
            if not filepath:
                if not webview:
                    return {"status": "error", "message": "Application window not available"}
                window = webview.active_window()
                if not window and len(webview.windows) > 0:
                    window = webview.windows[0]
                if not window:
                    return {"status": "error", "message": "Application window not available"}
                file_types = ('Image Files (*.png;*.jpg;*.jpeg;*.webp;*.bmp;*.gif)', 'All files (*.*)')
                result = window.create_file_dialog(webview.OPEN_DIALOG, allow_multiple=False, file_types=file_types)
                if not result:
                    return {"status": "cancelled"}
                filepath = result[0]
            mime_type, _ = mimetypes.guess_type(filepath)
            if not mime_type or not mime_type.startswith('image/'):
                mime_type = 'image/png'
            with open(filepath, 'rb') as img_file:
                b64_data = base64.b64encode(img_file.read()).decode('utf-8')
            data_url = f"data:{mime_type};base64,{b64_data}"
            self.db.update_faculty_photo(faculty_id, data_url)
            return {"status": "success", "profile_photo": data_url}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def save_faculty_photo_data(self, faculty_id, photo_data_url):
        try:
            self.db.update_faculty_photo(faculty_id, photo_data_url or "")
            return {"status": "success", "profile_photo": photo_data_url or ""}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def _hash_password(self, password):
        import hashlib
        return hashlib.sha256(password.encode()).hexdigest()

    def add_faculty(self, name, subject, password):
        try:
            pwd_hash = self._hash_password(password.strip())
            new_id = self.db.add_faculty(name, subject, pwd_hash)
            return {"status": "success", "id": new_id}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def verify_login(self, faculty_id, password):
        try:
            stored_hash = self.db.get_faculty_password(faculty_id)
            if stored_hash is None:
                return {"status": "error", "message": "Profile not found."}
            
            clean_pwd = (password or "").strip()
            # Check exact match, stripped match, or demo password ('password123')
            if (
                self._hash_password(password or "") == stored_hash
                or self._hash_password(clean_pwd) == stored_hash
                or clean_pwd == "password123"
            ):
                return {"status": "success"}
                
            return {"status": "error", "message": "Incorrect password."}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def add_student(self, name, roll_number, faculty_id=None):
        try:
            success = self.db.add_student(name, roll_number, faculty_id)
            if success:
                return {"status": "success"}
            else:
                return {"status": "error", "message": "Roll number already exists in this course!"}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def update_student_info(self, student_id, name, roll_number):
        try:
            ok, msg = self.db.update_student_info(student_id, name, roll_number)
            if ok:
                return {"status": "success", "message": msg}
            return {"status": "error", "message": msg}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def restore_student(self, student_id):
        try:
            self.db.restore_student(student_id)
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def bulk_import_students(self, faculty_id=None, filepath=None):
        try:
            import csv
            if not filepath:
                if not webview:
                    return {"status": "error", "message": "Window not available"}
                window = webview.active_window()
                if not window and len(webview.windows) > 0:
                    window = webview.windows[0]
                if not window:
                    return {"status": "error", "message": "Window not available"}
                file_types = ('CSV Files (*.csv)',)
                result = window.create_file_dialog(webview.OPEN_DIALOG, allow_multiple=False, file_types=file_types)
                if not result:
                    return {"status": "cancelled"}
                filepath = result[0]

            added_count = 0
            skipped_count = 0
            header_keywords = {"name", "student name", "student_name", "full name", "roll", "roll number", "roll_number", "roll no", "id", "sr no"}
            with open(filepath, 'r', encoding='utf-8-sig') as f:
                reader = csv.reader(f)
                for row_idx, row in enumerate(reader):
                    if not row or len(row) < 2:
                        continue
                    col0 = str(row[0] or "").strip()
                    col1 = str(row[1] or "").strip()
                    if not col0 or not col1:
                        continue
                    # Auto-detect whether the first row is a column header
                    if row_idx == 0 and (col0.lower() in header_keywords or col1.lower() in header_keywords):
                        continue
                    att_val = str(row[2]).strip() if len(row) >= 3 and str(row[2]).strip() else None
                    if self.db.add_student(col0, col1, faculty_id, attendance=att_val):
                        added_count += 1
                    else:
                        skipped_count += 1
                                
            return {"status": "success", "count": added_count, "skipped": skipped_count}
        except Exception as e:
            return {"status": "error", "message": f"Failed to import: {str(e)}"}

    def download_sample_csv(self):
        try:
            desktop = self._get_export_dir()
            target_path = os.path.join(desktop, "EduPulse_Student_Import_Sample.csv")
            DocumentExporter.export_sample_roster_csv(target_path)
            return {
                "status": "success",
                "path": target_path,
                "filename": os.path.basename(target_path),
                "download_url": f"/api/download/{os.path.basename(target_path)}"
            }
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def export_gradebook_csv(self, faculty_id=None):
        try:
            from datetime import datetime
            faculty_details = self.db.get_faculty_details(faculty_id) if faculty_id else None
            course_code = "".join([c if c.isalnum() else "_" for c in str((faculty_details or {}).get("course_code", "CS301"))])
            students = self.db.get_students_with_metrics(faculty_id)
            assignments = self.get_assignments(faculty_id)

            desktop = self._get_export_dir()
            target_path = os.path.join(desktop, f"EduPulse_Gradebook_{course_code}.csv")
            if os.path.exists(target_path):
                try:
                    with open(target_path, 'a'):
                        pass
                except (IOError, PermissionError):
                    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                    target_path = os.path.join(desktop, f"EduPulse_Gradebook_{course_code}_{ts}.csv")

            DocumentExporter.export_gradebook_csv(students, assignments, target_path)
            return {
                "status": "success",
                "path": target_path,
                "filename": os.path.basename(target_path),
                "download_url": f"/api/download/{os.path.basename(target_path)}"
            }
        except Exception as e:
            traceback.print_exc()
            return {"status": "error", "message": str(e)}

    def export_student_report_card(self, student_id, faculty_id=None):
        try:
            from datetime import datetime
            faculty_details = self.db.get_faculty_details(faculty_id) if faculty_id else None
            if not faculty_details:
                facs = self.db.get_all_faculties()
                first_id = facs[0][0] if facs else 1
                faculty_details = self.db.get_faculty_details(first_id) or {"name": "Faculty", "subject": "Data Science", "course_code": "CS301"}

            analytics = self.db.get_student_progress_analytics(faculty_id)
            students_list = analytics.get("students", [])
            student_data = next((s for s in students_list if int(s.get("id")) == int(student_id)), None)
            if not student_data:
                return {"status": "error", "message": "Student record not found."}

            desktop = self._get_export_dir()
            safe_roll = "".join([c if c.isalnum() or c in "-_" else "_" for c in str(student_data.get("roll", "Roll"))])
            safe_name = "".join([c if c.isalnum() or c in "-_" else "_" for c in str(student_data.get("name", "Student"))])
            target_path = os.path.join(desktop, f"EduPulse_ReportCard_{safe_roll}_{safe_name}.pdf")

            if os.path.exists(target_path):
                try:
                    with open(target_path, 'a'):
                        pass
                except (IOError, PermissionError):
                    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                    target_path = os.path.join(desktop, f"EduPulse_ReportCard_{safe_roll}_{safe_name}_{ts}.pdf")

            DocumentExporter.export_student_report_card_pdf(faculty_details, student_data, target_path)

            return {
                "status": "success",
                "path": target_path,
                "filename": os.path.basename(target_path),
                "download_url": f"/api/download/{os.path.basename(target_path)}"
            }
        except Exception as e:
            traceback.print_exc()
            return {"status": "error", "message": str(e)}

    def save_question_paper(self, faculty_id, title, chapter, difficulty, questions):
        try:
            paper_id = self.db.save_question_paper(faculty_id, title, chapter, difficulty, questions)
            return {"status": "success", "id": paper_id}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def get_question_papers(self, faculty_id=None):
        try:
            papers = self.db.get_question_papers(faculty_id)
            return {"status": "success", "papers": papers}
        except Exception as e:
            return {"status": "error", "message": str(e), "papers": []}

    def delete_question_paper(self, paper_id):
        try:
            self.db.delete_question_paper(paper_id)
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def delete_student(self, student_id, remark):
        try:
            self.db.delete_student(student_id, remark)
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def get_past_students(self, faculty_id=None):
        try:
            students = self.db.get_past_students(faculty_id)
            return [{"id": s[0], "name": s[1], "roll": s[2], "remark": s[3]} for s in students]
        except Exception as e:
            return []

    def get_students(self, faculty_id=None):
        try:
            return self.db.get_students_with_metrics(faculty_id)
        except Exception as e:
            print(f"Error fetching students with metrics: {e}")
            students = self.db.get_students(faculty_id)
            return [{"id": s[0], "name": s[1], "roll": s[2], "attendance": "--", "prev": "--", "latest": "--", "trend": "--", "status": "--", "insight": "--"} for s in students]

    def save_student_record(self, student_id, attendance, status, score=None, assignment_title="Assessment", custom_insight="", faculty_id=None, max_marks=100):
        try:
            self.db.update_student_record(student_id, attendance, status, score, assignment_title, custom_insight, faculty_id, max_marks)
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def batch_record_marks(self, assignment_title, records, faculty_id=None, assignment_id=None, max_marks=100):
        try:
            self.db.batch_record_marks(assignment_title, records, faculty_id, assignment_id, max_marks)
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def delete_assignment(self, assignment_id):
        try:
            self.db.delete_assignment(assignment_id)
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def get_assignment_marks(self, assignment_id):
        try:
            details = self.db.get_assignment_full_details(assignment_id)
            return {
                "status": "success",
                "marks": details["marks"],
                "percentages": details["percentages"],
                "statuses": details["statuses"],
                "max_marks": details["max_marks"]
            }
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def get_assignments(self, faculty_id=None):
        try:
            raw = self.db.get_assignments(faculty_id)
            return [
                {
                    "id": r[0],
                    "title": r[1],
                    "created_at": r[2],
                    "max_marks": float(r[3]) if len(r) > 3 and r[3] else 100.0
                }
                for r in raw
            ]
        except Exception as e:
            return []

    def update_student_attendance(self, student_id, attendance):
        try:
            self.db.update_student_attendance(student_id, attendance)
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def batch_update_attendance(self, records):
        try:
            self.db.batch_update_attendance(records)
            return {"status": "success"}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def get_progress_analytics(self, faculty_id=None):
        try:
            data = self.db.get_student_progress_analytics(faculty_id)
            return {"status": "success", "data": data}
        except Exception as e:
            import traceback
            traceback.print_exc()
            return {"status": "error", "message": str(e)}

    def backup_database(self):
        try:
            import sqlite3
            from contextlib import closing
            from datetime import datetime
            desktop = self._get_export_dir()
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            target_path = os.path.join(desktop, f"EduPulse_Backup_{ts}.db")

            with closing(sqlite3.connect(self.db.db_path)) as src_conn:
                with closing(sqlite3.connect(target_path)) as dst_conn:
                    src_conn.backup(dst_conn)

            return {
                "status": "success",
                "path": target_path,
                "filename": os.path.basename(target_path),
                "download_url": f"/api/download/{os.path.basename(target_path)}"
            }
        except Exception as e:
            traceback.print_exc()
            return {"status": "error", "message": str(e)}

    def restore_database(self, filepath=None):
        try:
            import sqlite3
            from contextlib import closing
            source_path = filepath
            if not source_path:
                if not webview:
                    return {"status": "error", "message": "Application window not available"}
                window = webview.active_window()
                if not window and len(webview.windows) > 0:
                    window = webview.windows[0]
                if not window:
                    return {"status": "error", "message": "Application window not available"}
                file_types = ('SQLite Database Backup (*.db;*.sqlite;*.sqlite3)', 'All files (*.*)')
                result = window.create_file_dialog(webview.OPEN_DIALOG, allow_multiple=False, file_types=file_types)
                if not result:
                    return {"status": "cancelled"}
                source_path = result[0]

            if not source_path or not os.path.exists(source_path):
                return {"status": "error", "message": "Selected backup file does not exist."}

            # Validate that source_path is a valid EduPulse SQLite database
            with closing(sqlite3.connect(source_path)) as check_conn:
                cur = check_conn.cursor()
                cur.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='Faculties'")
                if not cur.fetchone():
                    return {"status": "error", "message": "Invalid backup file: missing EduPulse tables."}

            with closing(sqlite3.connect(source_path)) as src_conn:
                with closing(sqlite3.connect(self.db.db_path)) as dst_conn:
                    src_conn.backup(dst_conn)

            # Re-run schema migration / backfill on restored DB to ensure full compatibility
            self.db.init_db()
            return {
                "status": "success",
                "filename": os.path.basename(source_path),
                "message": "Database restored successfully!"
            }
        except Exception as e:
            traceback.print_exc()
            return {"status": "error", "message": str(e)}



